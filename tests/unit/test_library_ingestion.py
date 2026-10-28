"""Library management drives the single ingestion pipeline: async upload, status in the library view, versions,
deletion, user retrieval limited to released practices, and the monitoring API."""

import io
import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from src.app.main import app as fastapi_app
from src.core.config import settings
from src.core.db.models import ChunkArchive, DocumentChunkLink, User
from src.core.db.session import session_scope
from src.workers.pipeline import LIBRARY_SCOPE
from tests.unit import _pipeline_fakes
from tests.unit.test_admin_library_pipeline import DOCX, _source

UPLOAD_KEYS = {"job_id", "status", "filename", "file_size_bytes", "sha256", "timestamp", "message"}
STATUS_KEYS = {
    "job_id", "status", "is_terminal", "practice_id", "filename", "stages", "current_stage_index", "progress_percent",
    "document_file_id", "counts", "chunk_count", "created_at", "updated_at", "completed_at", "file_size_bytes",
    "sha256", "parse_mode", "text_length", "failed_stage", "error_message",
}


@pytest.fixture
def env(tmp_path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(settings, "workspace_root", str(tmp_path))
    monkeypatch.setattr(settings, "lp_library_local_path", "")
    fakes = _pipeline_fakes.install(monkeypatch)
    return {**fakes, "client": TestClient(fastapi_app)}


def _login(client: TestClient, *, admin: bool) -> dict[str, str]:
    email = f"lib-{uuid.uuid4().hex[:8]}@example.com"
    token = client.post("/api/auth/login", json={"email": email, "password": "secret123"}).json()["access_token"]
    with session_scope() as db:
        db.scalar(select(User).where(User.email == email)).is_admin = admin
    return {"Authorization": f"Bearer {token}"}


def _upload(client, headers, name=None, content=None, **form):
    # Unique names by default: a file named like an existing practice's latest version becomes its next version.
    name = name or f"vendor-{uuid.uuid4().hex[:8]}.docx"
    data = {"owner": "Maya Patel", "rights": "Internal reuse permitted", "category": "Procurement", **form}
    content = content or _source(f"Vendor {uuid.uuid4().hex[:6]}", attributed=True)
    return client.post("/api/admin/library/uploads", headers=headers, data=data, files={"files": (name, content, DOCX)})


def _long_source(topic: str, sections: int) -> bytes:
    """A source long enough for several passages; more sections = the same text with sections appended."""
    doc = Document()
    doc.add_heading(f"{topic} handbook", 0)
    for n in range(1, sections + 1):
        doc.add_heading(f"{n} {topic} step {n}", 1)
        doc.add_paragraph(" ".join(f"Step {n} of {topic.lower()} records owner, evidence and decision {i}." for i in range(60)))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_upload_sources_is_asynchronous_and_runs_the_pipeline(env):
    client, headers = env["client"], _login(env["client"], admin=True)
    resp = _upload(client, headers, name="vendor-policy.docx")
    assert resp.status_code == 202, resp.text
    entry = resp.json()["jobs"][0]
    assert set(entry["ingestion"]) == UPLOAD_KEYS and entry["ingestion"]["filename"] == "vendor-policy.docx"
    practice = entry["practiceId"]  # the eager test worker has already run intake
    assert practice and entry["stage"] == "Quality review"

    status = client.get(f"/api/admin/library/documents/status/{entry['ingestion']['job_id']}", headers=headers).json()
    assert set(status) == STATUS_KEYS and status["status"] == "completed" and status["practice_id"] == practice
    assert status["counts"]["chunks"] == status["counts"]["embeddings"] > 0

    view = client.get(f"/api/admin/library?practice={practice}", headers=headers).json()
    assert view["ingestion"] == status  # selected practice: the same object as the status endpoint
    row = next(r for r in view["practiceInventory"] if r["id"] == practice)
    assert row["ingestion"] == status and row["indexing"] == "Indexed"
    assert view["pendingUploads"] == []
    details = client.get(f"/api/admin/library/documents/{status['job_id']}/details", headers=headers).json()
    assert details["postgres"]["chunk_count"] == details["qdrant"]["points_for_document"] > 0


def test_a_new_version_stays_out_of_user_search_and_keeps_the_old_one(env, monkeypatch):
    from src.query import pipeline as query_pipeline

    monkeypatch.setattr(query_pipeline, "claude_generate", lambda **_: "INSUFFICIENT_EVIDENCE")
    client, headers = env["client"], _login(env["client"], admin=True)
    first = _upload(client, headers).json()["jobs"][0]
    second = _upload(client, headers, name="vendor-v2.docx", practice_id=first["practiceId"]).json()["jobs"][0]
    status_1 = client.get(f"/api/admin/library/documents/status/{first['ingestion']['job_id']}", headers=headers).json()
    status_2 = client.get(f"/api/admin/library/documents/status/{second['ingestion']['job_id']}", headers=headers).json()
    store = env["get_store"]()
    # both versions stay indexed until a release; neither is a released practice, so users get nothing
    assert store.count(document_id=status_1["document_file_id"]) > 0 and store.count(document_id=status_2["document_file_id"]) > 0
    resp = client.post("/api/query", headers=headers, json={"question": "How are vendors screened?"})
    assert resp.status_code == 200, resp.text
    answer = resp.json()
    candidate_docs = {status_1["document_file_id"], status_2["document_file_id"]}
    assert not any(s["id"].rsplit("_c", 1)[0] in candidate_docs for s in answer["sources"])


def test_deleting_a_practice_removes_its_search_data(env):
    client, headers = env["client"], _login(env["client"], admin=True)
    entry = _upload(client, headers).json()["jobs"][0]
    status = client.get(f"/api/admin/library/documents/status/{entry['ingestion']['job_id']}", headers=headers).json()
    deleted = client.delete(f"/api/admin/library/{entry['practiceId']}", headers=headers).json()
    assert deleted["search_documents_removed"] == 1
    assert env["get_store"]().count(document_id=status["document_file_id"]) == 0
    jobs = client.get("/api/admin/library/documents/jobs", headers=headers).json()
    assert entry["ingestion"]["job_id"] not in [j["job_id"] for j in jobs]


def test_same_file_name_is_a_new_version_and_known_passages_are_not_stored_again(env):
    client, headers = env["client"], _login(env["client"], admin=True)
    topic, name = f"Sourcing {uuid.uuid4().hex[:6]}", f"sourcing-{uuid.uuid4().hex[:6]}.docx"
    original = _long_source(topic, 4)  # one build: a .docx records its save time
    first = _upload(client, headers, name=name, content=original).json()["jobs"][0]

    same = _upload(client, headers, name=name.upper(), content=original)
    assert same.status_code == 202
    unchanged = same.json()["jobs"][0]
    assert unchanged["outcome"] == "Unchanged" and unchanged["id"] == first["id"]
    assert unchanged["ingestion"]["status"] == "unchanged" and unchanged["ingestion"]["job_id"] == first["ingestion"]["job_id"]

    second = _upload(client, headers, name=name, content=_long_source(topic, 6)).json()["jobs"][0]
    assert second["practiceId"] == first["practiceId"] and second["id"] != first["id"]  # next version, same practice
    s1 = client.get(f"/api/admin/library/documents/status/{first['ingestion']['job_id']}", headers=headers).json()
    s2 = client.get(f"/api/admin/library/documents/status/{second['ingestion']['job_id']}", headers=headers).json()
    assert s2["status"] == "completed" and s2["counts"]["chunks"] > s1["counts"]["chunks"]
    details = client.get(f"/api/admin/library/documents/{s2['job_id']}/details", headers=headers).json()
    assert 0 < details["postgres"]["reused_chunk_count"] < details["postgres"]["chunk_count"]
    with session_scope() as db:
        shared = {c for (c,) in db.query(DocumentChunkLink.chunk_id).filter(DocumentChunkLink.document_file_id == s1["document_file_id"])}
        newer = {c for (c,) in db.query(DocumentChunkLink.chunk_id).filter(DocumentChunkLink.document_file_id == s2["document_file_id"])}
        assert shared & newer and db.query(ChunkArchive).filter(ChunkArchive.id.in_(shared | newer)).count() == len(shared | newer)
    store = env["get_store"]()
    assert store.count(document_id=s2["document_file_id"]) == len(newer)  # shared points list both versions

    # Deleting the practice removes both versions; shared passages go once nothing links to them.
    client.delete(f"/api/admin/library/{first['practiceId']}", headers=headers)
    with session_scope() as db:
        assert db.query(ChunkArchive).filter(ChunkArchive.id.in_(shared | newer)).count() == 0
    assert store.count(document_id=s1["document_file_id"]) == store.count(document_id=s2["document_file_id"]) == 0


def test_identical_content_under_another_name_is_rejected(env):
    client, headers = env["client"], _login(env["client"], admin=True)
    content = _source(f"Vendor {uuid.uuid4().hex[:6]}", attributed=True)
    assert _upload(client, headers, content=content).status_code == 202
    assert _upload(client, headers, content=content).status_code == 409


def test_monitoring_api_is_admin_only_and_old_entry_points_are_gone(env, monkeypatch):
    monkeypatch.setattr(settings, "advise_workbench_env", "production")
    client, member = env["client"], _login(env["client"], admin=False)
    assert client.get("/api/admin/library/documents/jobs", headers=member).status_code == 403
    admin = _login(client, admin=True)
    assert client.post("/api/admin/library/documents/upload-async", headers=admin,
                       files={"file": ("x.md", b"# x", "text/markdown")}).status_code in (404, 405)


def test_engagement_uploads_do_not_queue_the_pipeline():
    from src.app.api import documents

    assert not hasattr(documents, "queue_document")
    assert LIBRARY_SCOPE == "library"
