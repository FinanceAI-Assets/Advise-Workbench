"""The admin library pipeline: an uploaded source becomes a practice in the Leading Practice Library only after
its quality gates pass, an approver approves it, and an administrator publishes it."""

import io
import uuid
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from src.app.main import app as fastapi_app
from src.core.config import settings
from src.core.db.models import User
from src.core.db.session import session_scope
from src.memory.knowledge import lp_index, practice_ingest

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def _workspace_in_tmp(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "workspace_root", str(tmp_path))
    monkeypatch.setattr(settings, "lp_library_local_path", "")
    # Uploads run through the ingestion pipeline; run it in-process against local stand-ins.
    from tests.unit import _pipeline_fakes

    _pipeline_fakes.install(monkeypatch)


def _login(client: TestClient, prefix: str, *, admin: bool = False) -> tuple[dict[str, str], str]:
    email = f"{prefix}-{uuid.uuid4().hex[:8]}@example.com"
    resp = client.post("/api/auth/login", json={"email": email, "password": "secret123"})
    assert resp.status_code == 200
    with session_scope() as db:
        user = db.scalar(select(User).where(User.email == email))
        user.is_admin = admin
        user_id = user.id
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}, user_id


def _source(topic: str, *, attributed: bool) -> bytes:
    doc = Document()
    doc.add_heading(f"{topic} playbook", 0)
    doc.add_paragraph(f"This playbook sets out a repeatable method for {topic.lower()} so that risk, terms and service "
                      "levels are agreed before work starts.")
    doc.add_heading(f"1 {topic} risk screening", 1)
    doc.add_paragraph("Screen every counterparty for financial stability, data protection controls and sanctions "
                      f"exposure before {topic.lower()} begins, and record the owner and evidence of each check [1].")
    doc.add_heading("2 Commercial terms", 1)
    doc.add_paragraph(f"Agree payment terms and exit provisions for {topic.lower()} in one term sheet. Benchmarks "
                      "suggest a typical saving of 8 to 12 percent from structured negotiation [2].")
    doc.add_heading("References", 1)
    doc.add_paragraph("[1] Institute of Supply Management, Risk guide, 2024, https://example.org/risk")
    doc.add_paragraph("[2] Category benchmark study, 2023" if attributed else "[2] Industry benchmark (unattributed)")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, headers: dict[str, str], content: bytes, name: str, **form: str):
    data = {"owner": "Maya Patel", "rights": "Advisory Methods · Internal reuse permitted",
            "category": "Procurement", "maturity": "Defined", **form}
    return client.post("/api/admin/library/uploads", headers=headers, data=data, files={"files": (name, content, DOCX)})


def test_uploaded_practice_reaches_the_library_only_after_approval_and_publication() -> None:
    client = TestClient(fastapi_app)
    admin, _ = _login(client, "admin", admin=True)
    approver, approver_id = _login(client, "approver")
    topic = f"Onboarding {uuid.uuid4().hex[:6]}"

    # 1. Upload: the ingestion pipeline extracts, chunks, indexes and validates; reference 2 has no attributable source.
    source = _source(topic, attributed=False)
    upload = _upload(client, admin, source, "onboarding.docx")
    assert upload.status_code == 202, upload.text
    job = upload.json()["jobs"][0]
    practice_id = job["practiceId"]
    assert job["readable"] and job["stage"] == "Quality review"
    assert _upload(client, admin, source, "again.docx").status_code == 409  # the same file twice

    screen = client.get(f"/api/admin/library?practice={practice_id}", headers=admin).json()
    assert screen["selectedPracticeId"] == practice_id
    chunks = screen["extraction"]["chunks"]
    assert chunks and all(c["section"] for c in chunks)
    assert any(c["quality"] == "Reference 2 missing" for c in chunks)
    checks = {c["label"]: c for c in screen["indexValidation"]["checks"]}
    assert checks["Citation completeness"]["detail"] == "1 / 2 references attributed"
    assert checks["Index coverage"]["passed"] and checks["Locator integrity"]["passed"]
    assert screen["indexValidation"]["retrievalChecks"][0]["detail"].startswith("Relevant")
    gates = screen["qualityGates"]
    assert not gates["canSubmit"] and gates["blocker"]["title"] == "Unattributed source in reference 2"

    # Not in the library yet, and it cannot be submitted while a gate is open.
    library = client.get("/api/library", headers=approver).json()
    assert practice_id not in [c["id"] for c in library["cards"]]
    blocked = client.post(f"/api/admin/library/{practice_id}/submit", headers=admin, json={"approverId": approver_id})
    assert blocked.status_code == 409

    # 2. Correction: request it, then upload the corrected source against the same practice.
    correction = client.post(f"/api/admin/library/{practice_id}/request-correction", headers=admin,
                             json={"reason": "Attribute the benchmark in reference 2.", "assignee": "Maya Patel"})
    assert correction.status_code == 200
    corrected = _upload(client, admin, _source(topic, attributed=True), "onboarding-corrected.docx",
                        practice_id=practice_id, category="", maturity="")  # keeps the candidate's taxonomy
    assert corrected.status_code == 202, corrected.text
    screen = client.get(f"/api/admin/library?practice={practice_id}", headers=admin).json()
    assert all(g["passed"] for g in screen["qualityGates"]["gates"])
    assert screen["qualityGates"]["canSubmit"] and not screen["release"]["canPublish"]
    assert screen["release"]["gates"][0]["label"].startswith("Correction CR-")

    # 3. Approval: submit to an approver, who decides on the review page; only then can it be published.
    submit = client.post(f"/api/admin/library/{practice_id}/submit", headers=admin, json={"approverId": approver_id})
    assert submit.status_code == 200, submit.text
    review_id = submit.json()["reviewId"]
    early = client.post(f"/api/admin/library/{practice_id}/publish", headers=admin, json={"reason": "Release it now"})
    assert early.status_code == 409
    review = client.get(f"/api/reviews/{review_id}", headers=approver).json()
    assert review["canDecide"] and review["evidenceChecks"]
    decided = client.post(f"/api/reviews/{review_id}/decision", headers=approver, json={"decision": "approve"})
    assert decided.json()["state"] == "Approved"
    assert client.get(f"/api/admin/library?practice={practice_id}", headers=admin).json()["release"]["canPublish"]

    # 4. Publish: the practice is listed in the Leading Practice Library and served by its retrieval index.
    published = client.post(f"/api/admin/library/{practice_id}/publish", headers=admin,
                            json={"reason": "Approved for the autumn catalogue"})
    assert published.status_code == 200, published.text
    cards = client.get("/api/library", headers=approver, params={"q": topic}).json()["cards"]
    assert [c["id"] for c in cards] == [practice_id]
    assert cards[0]["category"] == "Procurement"
    detail = client.get(f"/api/library/{practice_id}", headers=approver).json()
    assert detail["status"] == "Approved" and detail["versionLabel"] == "v1.0"
    assert detail["coreGuidance"] and detail["practiceOwner"] == "Maya Patel"
    assert detail["sourceQuality"].startswith("2 / 2 citations verified")
    served = Path(settings.workspace_root) / "leading_practices" / "published" / f"{practice_id}.md"
    assert topic in served.read_text(encoding="utf-8")
    index = lp_index.load_index(force=True)
    assert any(c["source"] == "published" and "risk screening" in c["heading"].lower() for c in index.chunks)

    release = client.get(f"/api/admin/library?practice={practice_id}", headers=admin).json()["release"]
    assert release["published"] and release["notice"]["title"] == "v1.0 published"


def test_a_new_version_keeps_the_live_practice_served_until_it_is_released() -> None:
    client = TestClient(fastapi_app)
    admin, admin_id = _login(client, "admin", admin=True)
    member, _ = _login(client, "member")
    topic = f"Renewal {uuid.uuid4().hex[:6]}"
    practice_id = _upload(client, admin, _source(topic, attributed=True), "renewal.docx").json()["jobs"][0]["practiceId"]
    assert client.post(f"/api/admin/library/{practice_id}/submit", headers=admin, json={"approverId": admin_id}).status_code == 200
    client.post(f"/api/reviews/{practice_id.upper()}/decision", headers=admin, json={"decision": "approve"})
    assert client.post(f"/api/admin/library/{practice_id}/publish", headers=admin, json={"reason": "First release"}).status_code == 200

    second = _upload(client, admin, _source(f"{topic} v2", attributed=False), "renewal-v2.docx", practice_id=practice_id)
    assert second.status_code == 202
    live = client.get(f"/api/library/{practice_id}", headers=member).json()
    assert live["status"] == "Approved" and live["versionLabel"] == "v1.0"
    client.post(f"/api/admin/library/{practice_id}/request-correction", headers=admin, json={"reason": "Fix reference 2."})
    assert client.get(f"/api/library/{practice_id}", headers=member).json()["status"] == "Approved"
    row = next(r for r in client.get("/api/admin/library", headers=admin).json()["practiceInventory"] if r["id"] == practice_id)
    assert row["version"] == "v1.1 candidate" and row["lifecycle"] == "Candidate in review"


def test_unreadable_and_unsupported_sources_are_rejected_or_excluded() -> None:
    client = TestClient(fastapi_app)
    admin, _ = _login(client, "admin", admin=True)
    assert _upload(client, admin, b"MZ\x90\x00", "tool.exe").status_code == 422
    assert _upload(client, admin, b"not a zip", "broken.docx").status_code == 422
    empty = Document()
    buffer = io.BytesIO()
    empty.save(buffer)
    job = _upload(client, admin, buffer.getvalue(), f"empty-{uuid.uuid4().hex[:6]}.docx").json()["jobs"][0]
    assert not job["readable"] and job["outcome"] == "No readable text"
    screen = client.get(f"/api/admin/library?practice={job['practiceId']}", headers=admin).json()
    assert screen["extraction"]["error"]["title"] == "No readable text"
    assert not screen["qualityGates"]["canSubmit"]


def _deck(topic: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])  # title and content
    slide.shapes.title.text = f"{topic} operating model"
    slide.placeholders[1].text_frame.text = f"Assign one accountable owner for every {topic.lower()} decision."
    slide.notes_slide.notes_text_frame.text = "Presenter note: confirm owners with the sponsor before rollout."
    table_slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    table_slide.shapes.title.text = "Decision rights"
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for (r, c), text in {(0, 0): "Decision", (0, 1): "Owner", (1, 0): "Budget release", (1, 1): "Finance lead"}.items():
        table.cell(r, c).text = text
    group = table_slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)).text_frame.text = "Escalate unresolved items weekly."
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _workbook(topic: str) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Controls"
    sheet.append(["Control", "Owner", "Frequency"])
    sheet.append([f"{topic} access review", "IT risk", "Quarterly"])
    book.create_sheet("Metrics").append(["Cycle time", 12, None, "days"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_powerpoint_and_excel_sources_are_read_slide_by_slide_and_sheet_by_sheet() -> None:
    deck = practice_ingest.extract_blocks("model.pptx", _deck("Vendor"))
    assert deck["parseMode"] == "pptx" and deck["pages"] == 2 and deck["readable"]
    texts = [(b["kind"], b["text"], b["page"]) for b in deck["blocks"]]
    assert ("heading", "Vendor operating model", 1) in texts
    assert ("para", "Assign one accountable owner for every vendor decision.", 1) in texts
    assert ("para", "Presenter note: confirm owners with the sponsor before rollout.", 1) in texts
    assert ("para", "Budget release | Finance lead", 2) in texts
    assert ("para", "Escalate unresolved items weekly.", 2) in texts  # text inside a grouped shape
    assert sum(1 for kind, text, _ in texts if text == "Vendor operating model") == 1  # the title is not repeated

    book = practice_ingest.extract_blocks("controls.xlsx", _workbook("Vendor"))
    assert book["parseMode"] == "xlsx" and book["pages"] == 2 and book["readable"]
    rows = [(b["kind"], b["text"], b["page"]) for b in book["blocks"]]
    assert rows[0] == ("heading", "Controls", 1)
    assert ("para", "Vendor access review | IT risk | Quarterly", 1) in rows
    assert ("heading", "Metrics", 2) in rows
    assert ("para", "Cycle time | 12 | days", 2) in rows  # empty cells are skipped


def test_word_powerpoint_and_excel_sources_upload_together() -> None:
    client = TestClient(fastapi_app)
    admin, _ = _login(client, "admin", admin=True)
    topic = f"Mixed {uuid.uuid4().hex[:6]}"
    files = [
        ("files", (f"{topic}.docx", _source(topic, attributed=True), DOCX)),
        ("files", (f"{topic}.pptx", _deck(topic), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
        ("files", (f"{topic}.xlsx", _workbook(topic), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
    ]
    data = {"owner": "Maya Patel", "rights": "Advisory Methods · Internal reuse permitted"}
    resp = client.post("/api/admin/library/uploads", headers=admin, data=data, files=files)
    assert resp.status_code == 201, resp.text
    jobs = resp.json()["jobs"]
    assert len(jobs) == 3 and all(job["readable"] for job in jobs)
    assert len({job["practiceId"] for job in jobs}) == 3  # each file is its own candidate practice


def test_older_office_formats_and_broken_archives_are_rejected_with_a_reason() -> None:
    client = TestClient(fastapi_app)
    admin, _ = _login(client, "admin", admin=True)
    legacy = _upload(client, admin, b"\xd0\xcf\x11\xe0legacy", "deck.ppt")
    assert legacy.status_code == 422 and "Save it as .pptx" in legacy.json()["detail"]
    assert "Save it as .xlsx" in _upload(client, admin, b"\xd0\xcf\x11\xe0legacy", "sheet.xls").json()["detail"]
    broken = _upload(client, admin, b"not a zip", "broken.pptx")
    assert broken.status_code == 422 and "not a valid PPTX file" in broken.json()["detail"]


def test_member_cannot_upload_library_sources(monkeypatch: pytest.MonkeyPatch) -> None:
    # Development mode treats every signed-in user as an administrator (src/core/auth.py); outside it, it is enforced.
    monkeypatch.setattr(settings, "advise_workbench_env", "production")
    client = TestClient(fastapi_app)
    member, _ = _login(client, "member")
    assert _upload(client, member, _source("Any", attributed=True), "any.docx").status_code == 403


def test_chunker_respects_headings_and_carries_overlap() -> None:
    long_paragraph = " ".join(["Agree the baseline with the sponsor and record the evidence."] * 40)
    blocks = [
        {"kind": "heading", "text": "1 Baseline", "page": 1, "level": 1},
        {"kind": "para", "text": long_paragraph, "page": 1, "level": 0},
        {"kind": "para", "text": long_paragraph, "page": 2, "level": 0},
        {"kind": "heading", "text": "2 Reporting", "page": 3, "level": 1},
        {"kind": "para", "text": "Report monthly using the approved measures and the same baseline.", "page": 3, "level": 0},
    ]
    chunks = practice_ingest.chunk_blocks(blocks)
    assert [c["section"] for c in chunks] == ["§1 Baseline", "§1 Baseline", "§2 Reporting"]
    assert chunks[1]["page"] == 2 and chunks[2]["page"] == 3
    overlap = practice_ingest._overlap_tail(chunks[0]["text"])
    assert chunks[1]["text"].startswith(overlap)
    assert not chunks[2]["text"].startswith(overlap)
