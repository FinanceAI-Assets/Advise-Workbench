"""Tests for the workbench reviews, library, notifications, search and admin console APIs."""

import uuid

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from src.app.main import app as fastapi_app
from src.core.config import settings
from src.core.db.models import LpPractice, Review, ReviewComment, Run, User
from src.core.db.session import session_scope

BRIEF = "Design a service ownership pilot for shared services over ninety days."


@pytest.fixture(autouse=True)
def _workspace_in_tmp(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Creating an engagement scaffolds a workspace folder; keep those out of the repo's workspace."""
    monkeypatch.setattr(settings, "workspace_root", str(tmp_path))


def _email(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}@example.com"


def _login(client: TestClient, email: str) -> dict[str, str]:
    resp = client.post("/api/auth/login", json={"email": email, "password": "secret123"})
    assert resp.status_code == 200
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def _user_id(email: str) -> str:
    with session_scope() as db:
        return db.scalar(select(User.id).where(User.email == email))


def _engagement(client: TestClient, headers: dict[str, str]) -> str:
    resp = client.post("/api/engagements", json={"name": "Northstar pilot", "brief": BRIEF}, headers=headers)
    assert resp.status_code == 201
    return resp.json()["id"]


class Workspace:
    """An engagement with a lead and an independent reviewer, plus a finished run to review."""

    def __init__(self, client: TestClient) -> None:
        self.client = client
        self.lead_email, self.reviewer_email = _email("lead"), _email("reviewer")
        self.lead, self.reviewer = _login(client, self.lead_email), _login(client, self.reviewer_email)
        self.pid = _engagement(client, self.lead)
        added = client.post(
            f"/api/engagements/{self.pid}/members",
            json={"email": self.reviewer_email, "role": "Independent reviewer"},
            headers=self.lead,
        )
        assert added.status_code == 201
        self.run_id = f"run_{uuid.uuid4().hex[:8]}"
        with session_scope() as db:
            db.add(Run(id=self.run_id, project_id=self.pid, status="review_ready", output_types='["pptx"]',
                       instruction="Roadmap", plan_payload="{}"))

    def submit(self) -> str:
        resp = self.client.post(
            "/api/reviews",
            json={"projectId": self.pid, "runId": self.run_id, "artifactType": "pptx",
                  "title": "Transformation roadmap", "reviewerId": _user_id(self.reviewer_email),
                  "dueDate": "2026-10-12"},
            headers=self.lead,
        )
        assert resp.status_code == 201, resp.text
        return resp.json()["id"]


def test_author_can_review_and_approve_their_own_draft() -> None:
    """A solo engagement: the author submits to themselves, approves, and export unlocks."""
    client = TestClient(fastapi_app)
    ws = Workspace(client)
    resp = client.post(
        "/api/reviews",
        json={"projectId": ws.pid, "runId": ws.run_id, "artifactType": "pptx", "title": "Roadmap",
              "reviewerId": _user_id(ws.lead_email)},
        headers=ws.lead,
    )
    assert resp.status_code == 201
    rid = resp.json()["id"]
    detail = client.get(f"/api/reviews/{rid}", headers=ws.lead).json()
    assert detail["canDecide"] is True and detail["selfReview"] is True

    approved = client.post(f"/api/reviews/{rid}/decision", json={"decision": "approve"}, headers=ws.lead)
    assert approved.json()["state"] == "Approved"
    assert client.get(f"/api/engagements/{ws.pid}/export-readiness", headers=ws.lead).json()["ready"] is True
    # The self-approval is labelled as such for governance.
    timeline = client.get(f"/api/reviews/{rid}", headers=ws.lead).json()["approvalTimeline"]
    assert timeline[-1]["title"].startswith("Self-approved")
    # Someone who is not a member still cannot be the reviewer.
    outsider_email = _email("outsider")
    _login(client, outsider_email)
    with session_scope() as db:
        db.add(Run(id=ws.run_id + "b", project_id=ws.pid, status="review_ready", output_types='["docx"]',
                   instruction="Doc", plan_payload="{}"))
    refused = client.post(
        "/api/reviews",
        json={"projectId": ws.pid, "runId": ws.run_id + "b", "artifactType": "docx", "title": "Memo",
              "reviewerId": _user_id(outsider_email)},
        headers=ws.lead,
    )
    assert refused.status_code == 422


def test_review_round_trip_changes_then_approval() -> None:
    client = TestClient(fastapi_app)
    ws = Workspace(client)
    rid = ws.submit()
    assert client.get(f"/api/engagements/{ws.pid}", headers=ws.lead).json()["stage"] == "In review"

    # The owner cannot decide; the reviewer must say what to change.
    assert client.post(f"/api/reviews/{rid}/decision", json={"decision": "approve"}, headers=ws.lead).status_code == 403
    empty = client.post(f"/api/reviews/{rid}/decision", json={"decision": "request_changes"}, headers=ws.reviewer)
    assert empty.status_code == 422
    asked = client.post(
        f"/api/reviews/{rid}/decision",
        json={"decision": "request_changes", "note": "Name the pilot owner."},
        headers=ws.reviewer,
    )
    assert asked.json()["state"] == "Changes requested"

    detail = client.get(f"/api/reviews/{rid}", headers=ws.lead).json()
    assert detail["canRevise"] is True and detail["openCount"] == 1
    # A revision is refused while a comment is open.
    assert client.post(f"/api/reviews/{rid}/submit-revision", json={}, headers=ws.lead).status_code == 409
    comment_id = detail["comments"][0]["id"]
    assert client.post(f"/api/reviews/{rid}/comments/{comment_id}/resolve", headers=ws.lead).status_code == 200
    revised = client.post(f"/api/reviews/{rid}/submit-revision", json={}, headers=ws.lead).json()
    assert revised["state"] == "In review" and revised["versionLabel"] == "PPTX v0.2"

    approved = client.post(f"/api/reviews/{rid}/decision", json={"decision": "approve"}, headers=ws.reviewer)
    assert approved.json()["state"] == "Approved"
    gate = client.get(f"/api/engagements/{ws.pid}/export-readiness", headers=ws.lead).json()
    assert gate["ready"] is True
    assert client.get(f"/api/engagements/{ws.pid}", headers=ws.lead).json()["stage"] == "Approved"


def test_review_is_hidden_from_non_members_and_notifies_the_reviewer() -> None:
    client = TestClient(fastapi_app)
    ws = Workspace(client)
    rid = ws.submit()
    outsider = _login(client, _email("outsider"))
    assert client.get(f"/api/reviews/{rid}", headers=outsider).status_code == 403
    assert client.get("/api/reviews", headers=outsider).json()["total"] == 0

    inbox = client.get("/api/notifications?filter=unread", headers=ws.reviewer).json()
    assert inbox["counts"]["unread"] == 1 and inbox["items"][0]["link"] == f"/reviews/{rid}"
    assert client.post("/api/notifications/read-all", headers=ws.reviewer).status_code == 200
    assert client.get("/api/notifications", headers=ws.reviewer).json()["counts"]["unread"] == 0
    assert client.get("/api/notifications?filter=bogus", headers=ws.reviewer).status_code == 422


def _practice(status: str = "Approved") -> str:
    pid = f"lp-{uuid.uuid4().hex[:6]}"
    with session_scope() as db:
        db.add(LpPractice(id=pid, title=f"Ownership playbook {pid}", description="Service ownership method.",
                          category="Operating model", maturity="Managed", version_label="v1.0", status=status,
                          owner="Maya Patel"))
    return pid


def test_library_shows_approved_practices_and_saves_per_user() -> None:
    client = TestClient(fastapi_app)
    headers = _login(client, _email("reader"))
    approved, candidate = _practice(), _practice("Quality review")

    listing = client.get(f"/api/library?q={approved}", headers=headers).json()
    assert [c["id"] for c in listing["cards"]] == [approved]
    assert client.get(f"/api/library?q={candidate}", headers=headers).json()["cards"] == []

    assert client.post(f"/api/library/{approved}/save", headers=headers).json()["saved"] is True
    saved = client.get("/api/library?saved_only=true", headers=headers).json()
    assert [c["id"] for c in saved["cards"]] == [approved]
    assert client.delete(f"/api/library/{approved}/save", headers=headers).json()["saved"] is False
    assert client.get("/api/library/lp-missing", headers=headers).status_code == 404


def test_practice_added_to_engagement_counts_as_a_source_and_is_searchable() -> None:
    client = TestClient(fastapi_app)
    headers = _login(client, _email("lead"))
    pid = _engagement(client, headers)
    approved, candidate = _practice(), _practice("Quality review")

    refused = client.post(f"/api/library/{candidate}/add-to-engagement", json={"projectId": pid}, headers=headers)
    assert refused.status_code == 409
    added = client.post(f"/api/library/{approved}/add-to-engagement", json={"projectId": pid}, headers=headers)
    assert added.json()["added"] is True
    again = client.post(f"/api/library/{approved}/add-to-engagement", json={"projectId": pid}, headers=headers)
    assert again.json()["added"] is False
    assert client.get(f"/api/engagements/{pid}", headers=headers).json()["tabCounts"]["sources"] == 1

    found = client.get("/api/search?q=northstar", headers=headers).json()
    assert [e["id"] for e in found["engagements"]] == [pid]
    assert client.get("/api/search?q=n", headers=headers).json()["engagements"] == []


def test_conversations_can_be_started_and_switched() -> None:
    client = TestClient(fastapi_app)
    headers = _login(client, _email("lead"))
    pid = _engagement(client, headers)
    first = client.post(f"/api/engagements/{pid}/conversations", json={"title": "Roadmap planning"}, headers=headers)
    second = client.post(f"/api/engagements/{pid}/conversations", json={"title": "Service options"}, headers=headers)
    assert first.status_code == second.status_code == 201

    listing = client.get(f"/api/engagements/{pid}/conversations", headers=headers).json()
    assert [c["label"] for c in listing["items"]] == ["Service options", "Roadmap planning"]
    client.post(f"/api/engagements/{pid}/conversations/{first.json()['id']}/activate", headers=headers)
    listing = client.get(f"/api/engagements/{pid}/conversations", headers=headers).json()
    assert listing["items"][0]["label"] == "Roadmap planning" and listing["items"][0]["active"] is True
    # The existing conversation endpoint follows the active thread.
    active = client.get(f"/api/projects/{pid}/conversation", headers=headers).json()
    assert active["conversation_id"] == first.json()["id"]


def _make_admin(email: str) -> None:
    with session_scope() as db:
        db.scalar(select(User).where(User.email == email)).is_admin = True


def test_admin_console_requires_an_administrator() -> None:
    client = TestClient(fastapi_app)
    headers = _login(client, _email("member"))
    for path in ("/api/admin/overview", "/api/admin/engagements", "/api/admin/reviews", "/api/admin/export/audit.csv"):
        assert client.get(path, headers=headers).status_code == 403


def test_admin_routes_reviews_and_manages_retention_but_cannot_bypass_gates() -> None:
    client = TestClient(fastapi_app)
    ws = Workspace(client)
    rid = ws.submit()
    admin_email = _email("admin")
    admin = _login(client, admin_email)
    _make_admin(admin_email)

    oversight = client.get(f"/api/admin/engagements?selected={ws.pid}", headers=admin).json()
    assert oversight["selected"]["id"] == ws.pid and len(oversight["selected"]["members"]) == 2
    assert next(r for r in oversight["rows"] if r["id"] == ws.pid)["controlState"] == "Review pending"

    # Changes need a recorded reason.
    assert client.post(f"/api/admin/engagements/{ws.pid}/retention", json={"days": 365, "reason": ""},
                       headers=admin).status_code == 422
    kept = client.post(f"/api/admin/engagements/{ws.pid}/retention",
                       json={"days": 365, "reason": "Annual adoption evaluation"}, headers=admin)
    assert kept.json()["retentionDays"] == 365
    # An open review blocks archiving.
    blocked = client.post(f"/api/admin/engagements/{ws.pid}/archive", json={"reason": "Engagement closed"},
                          headers=admin)
    assert blocked.status_code == 409

    # The owner cannot be made the reviewer; another account can.
    self_review = client.post(f"/api/admin/reviews/{rid}/assign",
                              json={"reviewerId": _user_id(ws.lead_email), "reason": "Reviewer on leave"},
                              headers=admin)
    assert self_review.status_code == 422
    routed = client.post(f"/api/admin/reviews/{rid}/assign",
                         json={"reviewerId": _user_id(admin_email), "reason": "Reviewer on leave"}, headers=admin)
    assert routed.status_code == 200

    queue = client.get(f"/api/admin/reviews?selected={rid}", headers=admin).json()
    assert queue["selectedReview"]["id"] == rid
    assert any("assigned to" in e["title"] for e in queue["fullAuditTrail"])
    export = client.get("/api/admin/export/reviews.csv", headers=admin)
    assert export.status_code == 200 and rid in export.text


def test_practice_cannot_be_published_before_independent_approval() -> None:
    client = TestClient(fastapi_app)
    admin_email = _email("admin")
    admin = _login(client, admin_email)
    _make_admin(admin_email)
    practice = _practice("Quality review")
    with session_scope() as db:
        db.add(Review(id=practice.upper(), title="Ownership playbook", artifact_type="practice", version_label="v1.0",
                      state="Quality review"))
    reason = {"reason": "Release for the autumn catalogue"}

    assert client.post(f"/api/admin/library/{practice}/publish", json=reason, headers=admin).status_code == 409
    with session_scope() as db:
        db.scalar(select(Review).where(Review.id == practice.upper())).state = "Approved"
    assert client.post(f"/api/admin/library/{practice}/publish", json=reason, headers=admin).json()["status"] == "Approved"


def test_seed_is_idempotent_and_clears_only_demo_data() -> None:
    from scripts import seed_demo

    client = TestClient(fastapi_app)
    keeper = _login(client, _email("keeper"))
    kept = _engagement(client, keeper)

    seed_demo.seed("demo-pass-123")
    seed_demo.seed("demo-pass-123")
    alex = client.post("/api/auth/login", json={"email": seed_demo.demo_email("alex.chen"), "password": "demo-pass-123"})
    headers = {"Authorization": f"Bearer {alex.json()['access_token']}"}
    assert client.get("/api/engagements?page_size=50", headers=headers).json()["total"] == len(seed_demo.ENGAGEMENTS)
    stats = {s["id"]: s for s in client.get("/api/engagements/stats", headers=headers).json()["items"]}
    assert stats["in-review"]["helper"] == "2 need your response"
    review = client.get("/api/reviews/RV-104", headers=headers).json()
    assert review["canRevise"] is True and review["openCount"] == 2
    with session_scope() as db:
        assert db.scalars(select(ReviewComment).where(ReviewComment.review_id == "RV-104")).all().__len__() == 2
    artifacts = client.get(f"/api/runs/{seed_demo.NORTHSTAR}/{seed_demo.RUN_ID}/artifacts", headers=headers).json()
    assert len(artifacts["artifacts"]["pptx_slides"]) == 12
    docx = client.get(f"/api/runs/{seed_demo.NORTHSTAR}/{seed_demo.RUN_ID}/artifacts/docx/download", headers=headers)
    assert docx.status_code == 200

    jordan = client.post("/api/auth/login", json={"email": seed_demo.demo_email("jordan.kim"), "password": "demo-pass-123"})
    admin = {"Authorization": f"Bearer {jordan.json()['access_token']}"}
    assert client.get("/api/admin/library", headers=admin).status_code == 200
    blocked = client.post("/api/admin/export-requests/EX-031/approve", json={"reason": "Sponsor review"}, headers=admin)
    assert blocked.status_code == 409
    released = client.post("/api/admin/export-requests/EX-028/approve", json={"reason": "Sponsor review"}, headers=admin)
    assert released.json()["state"] == "Released"

    seed_demo.clear()
    assert client.get("/api/engagements", headers=keeper).json()["items"][0]["id"] == kept
    with session_scope() as db:
        assert db.scalar(select(User).where(User.email == seed_demo.demo_email("alex.chen"))) is None


def test_conversation_usage_is_tracked_live_and_saved_per_conversation() -> None:
    from src.app.api.projects._conversation_usage import live_key, record_turn_usage
    from src.llm.run_budget import capture_usage, charge_llm_usage

    client = TestClient(fastapi_app)
    email = _email("lead")
    headers = _login(client, email)
    pid = _engagement(client, headers)
    first = client.post(f"/api/engagements/{pid}/conversations", json={"title": "Roadmap planning"}, headers=headers)
    url = f"/api/engagements/{pid}/conversations/usage"
    empty = client.get(url, headers=headers).json()
    assert empty["totalTokens"] == 0 and empty["live"] is False and empty["model"]

    with capture_usage(live_key(pid, _user_id(email))) as usage:
        charge_llm_usage(input_tokens=1000, output_tokens=200)
        # While the reply is being written its tokens already show.
        live = client.get(url, headers=headers).json()
        assert live["live"] is True and live["inputTokens"] == 1000 and live["outputTokens"] == 200
    with session_scope() as db:
        record_turn_usage(db, conversation_id=first.json()["id"], project_id=pid, usage=usage.snapshot())

    saved = client.get(url, headers=headers).json()
    assert saved["live"] is False and saved["totalTokens"] == 1200 and saved["replies"] == 1
    assert saved["costUsd"] > 0
    # A new conversation starts its own count.
    client.post(f"/api/engagements/{pid}/conversations", json={"title": "Second thread"}, headers=headers)
    assert client.get(url, headers=headers).json()["totalTokens"] == 0


def test_run_progress_reports_milestones_from_run_events() -> None:
    from src.core.db.models import RunEvent

    client = TestClient(fastapi_app)
    headers = _login(client, _email("lead"))
    pid = _engagement(client, headers)
    url = f"/api/engagements/{pid}/run-progress"
    assert client.get(url, headers=headers).json() == {"run": None}

    run_id = f"run_{uuid.uuid4().hex[:8]}"
    with session_scope() as db:
        db.add(Run(id=run_id, project_id=pid, status="running", output_types='["pptx"]', instruction="Deck",
                   plan_payload="{}"))
        db.flush()
        db.add(RunEvent(run_id=run_id, event_type="step", payload='{"payload": {"status": "execution_started"}}'))
        db.add(RunEvent(run_id=run_id, event_type="fanout_start", payload="{}"))

    working = client.get(url, headers=headers).json()["run"]
    assert working["active"] is True and working["finished"] is False
    assert [s["state"] for s in working["stages"]] == ["done", "done", "current", "upcoming", "upcoming", "upcoming"]

    with session_scope() as db:
        db.scalar(select(Run).where(Run.id == run_id)).status = "review_ready"
    finished = client.get(url, headers=headers).json()["run"]
    assert finished["finished"] is True and all(s["state"] == "done" for s in finished["stages"])


def test_rendered_slides_keep_one_title_and_no_empty_placeholders() -> None:
    """Regression: the layout title placeholder and the composer's title box both carried the title."""
    from pptx import Presentation
    from pptx.util import Inches

    from src.deliverables.pptx.pptx_layouts import add_slide_with_layout, tidy_slide_placeholders

    prs = Presentation()
    title = "Approve the four-phase programme now"
    slide = add_slide_with_layout(prs, "bullets", title=title)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    box.text_frame.text = title
    tidy_slide_placeholders(slide)

    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert texts.count(title) == 1
    assert slide.shapes.title.text_frame.text == title
    assert slide.shapes.title.top == Inches(0.5)
    assert all(p.text_frame.text.strip() for p in slide.placeholders)


def test_agreed_slides_become_the_deck_outline() -> None:
    import json as _json

    from src.agents.conversation.slide_negotiator import agreed_deck_outline
    from src.core.db.models import MemoryItem, Project

    pid = f"p_{uuid.uuid4().hex[:10]}"
    with session_scope() as db:
        email = _email("lead")
        TestClient(fastapi_app).post("/api/auth/login", json={"email": email, "password": "secret123"})
        uid = _user_id(email)
        db.add(Project(id=pid, name="Outline check", created_by=uid))
        db.flush()
        assert agreed_deck_outline(db, project_id=pid) is None
        for n in (1, 2, 3):
            db.add(MemoryItem(id=f"mi_{uuid.uuid4().hex[:10]}", project_id=pid, memory_type="decision", key=f"slide_{n}",
                              value=_json.dumps({"slide_num": n, "title": f"Slide {n}", "slide_type": "column_cards",
                                                 "key_message": f"Message {n}"})))
        db.add(MemoryItem(id=f"mi_{uuid.uuid4().hex[:10]}", project_id=pid, memory_type="decision", key="storyline_arc",
                          value=_json.dumps({"arc": "case_led"})))
        db.flush()
        outline = agreed_deck_outline(db, project_id=pid)
    assert outline is not None and outline["narrative_arc"] == "case_led"
    assert [s["title"] for s in outline["slides"]] == ["Slide 1", "Slide 2", "Slide 3"]
