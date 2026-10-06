"""Slides written the way the agents actually write them must render with their content (found in a live run)."""

from __future__ import annotations

from pptx import Presentation

from src.deliverables.pptx.pptx_artifact_renderer import _normalize_slide_content, render_pptx_with_artifact_tool

TWO_BY_TWO = {
    "slide_type": "two_by_two",
    "title": "Three Pillars of P2P Risk",
    "quadrants": [
        {"heading": "Fragmentation", "accent": "green", "body": "14 steps across 6 roles diffuse accountability."},
        {"heading": "Manual AP", "accent": "dark", "body": "Invoice capture and matching are manual."},
        {"heading": "Controls", "accent": "gray", "body": "SOX checks happen late."},
        {"heading": "Systems", "accent": "dark", "body": "SAP ECC and NetSuite are not integrated."},
    ],
}
HARVEY = {
    "slide_type": "harvey_balls",
    "title": "Manual Process Drives Cycle Time",
    "rows": [
        {"label": "Procurement", "description": "3-5 days", "ball_values": ["empty", "half", "full"]},
        {"label": "Goods Receipt", "description": "2-3 days", "ball_values": ["quarter", "three quarter", 4]},
    ],
}
BULLETS = {"slide_type": "bullets", "title": "Pillars", "bullets": ["**Automation**: OCR and rules engine.", "__Control__: policy checks."]}


def _slide_text(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_field_variants_are_mapped() -> None:
    quads = _normalize_slide_content(TWO_BY_TWO)["quadrants"]
    assert [q["label"] for q in quads] == ["Fragmentation", "Manual AP", "Controls", "Systems"]
    rows = _normalize_slide_content(HARVEY)["rows"]
    assert [r["scores"] for r in rows] == [[0, 2, 4], [1, 3, 4]]
    assert _normalize_slide_content(BULLETS)["bullets"] == ["Automation: OCR and rules engine.", "Control: policy checks."]
    assert _normalize_slide_content({"bullets": ["5 ** 2 stays", "a**b"]})["bullets"] == ["5 ** 2 stays", "a**b"]
    # The model's own page counter in footer_note is dropped; a real note after it is kept.
    assert "footer_note" not in _normalize_slide_content({"footer_note": "01 / 10"})
    assert "footer_note" not in _normalize_slide_content({"footer_note": "CONFIDENTIAL · 05 / 10"})
    assert _normalize_slide_content({"footer_note": "CONFIDENTIAL · 06 / 10; Median cycle ~14 days"})["footer_note"] == "Median cycle ~14 days"
    assert _normalize_slide_content({"footer_note": "Source: AP memo"})["footer_note"] == "Source: AP memo"


def test_quadrants_and_bullets_render_with_their_text(tmp_path) -> None:
    result = render_pptx_with_artifact_tool({"pptx_slides": [TWO_BY_TWO, HARVEY, BULLETS]}, tmp_path, None)
    prs = Presentation(str(result["output_path"]))
    texts = [_slide_text(s) for s in prs.slides]
    deck = "\n".join(texts)
    for word in ("Fragmentation", "14 steps across 6 roles", "Systems", "Automation: OCR"):
        assert word in deck
    assert "**" not in deck and "__Control__" not in deck
    harvey_slide = next(s for s in prs.slides if "Manual Process Drives Cycle Time" in _slide_text(s))
    assert any(sh.shape_type == 13 for sh in harvey_slide.shapes)  # the figure picture is present
    assert "02 / 03" in _slide_text(harvey_slide)  # figure slides carry the same footer as other slides
